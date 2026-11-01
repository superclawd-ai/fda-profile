"""Build a clean 16:9 PowerPoint deck from a JSON outline (see SKILL.md).

usage: fda-python .claude/skills/make-slides/slides.py outline.json out.pptx [--template t.pptx] [--pdf]
"""
import argparse
import json
import os
import pathlib
import subprocess

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

ap = argparse.ArgumentParser()
ap.add_argument("outline")
ap.add_argument("out")
ap.add_argument("--template")
ap.add_argument("--pdf", action="store_true")
a = ap.parse_args()

src = pathlib.Path(a.outline).resolve()
o = json.loads(src.read_text(encoding="utf-8"))
base = src.parent
accent = RGBColor.from_string(o.get("accent", "2457D6"))
dark = RGBColor(0x1D, 0x1D, 0x1F)
grey = RGBColor(0x6B, 0x72, 0x80)
FONT = o.get("font", "Inter")

prs = Presentation(a.template) if a.template else Presentation()
if not a.template:
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
M = Inches(0.6)


def text(slide, x, y, w, h, s, size, bold=False, color=dark, align=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = s
    f = p.runs[0].font
    f.size, f.bold, f.name = Pt(size), bold, FONT
    f.color.rgb = color
    if align is not None:
        p.alignment = align
    return tf


def bar(slide):
    r = slide.shapes.add_shape(1, 0, 0, Inches(0.18), H)
    r.fill.solid()
    r.fill.fore_color.rgb = accent
    r.line.fill.background()


def title(slide, s):
    bar(slide)
    text(slide, M, Inches(0.45), W - 2 * M, Inches(1.0), s, 30, bold=True)


def bullets(slide, items, x, y, w, h, size=20):
    tf = slide.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"•  {it}"
        p.space_after = Pt(10)
        f = p.runs[0].font
        f.size, f.name = Pt(size), FONT
        f.color.rgb = dark


def notes(slide, s):
    if s:
        slide.notes_slide.notes_text_frame.text = s


# Title slide
s = prs.slides.add_slide(blank)
bar(s)
text(s, M, Inches(2.6), W - 2 * M, Inches(1.4), o.get("title", ""), 44, bold=True)
if o.get("subtitle"):
    text(s, M, Inches(3.9), W - 2 * M, Inches(0.8), o["subtitle"], 20, color=grey)

for sl in o.get("slides", []):
    s = prs.slides.add_slide(blank)
    if sl.get("type") == "section":
        bar(s)
        text(s, M, Inches(3.0), W - 2 * M, Inches(1.2), sl.get("title", ""), 38, bold=True, color=accent)
        notes(s, sl.get("notes"))
        continue
    title(s, sl.get("title", ""))
    top, avail_h = Inches(1.7), H - Inches(2.2)
    if "image" in sl:
        pic = s.shapes.add_picture(str(base / sl["image"]), M, top)
        scale = min((W - 2 * M) / pic.width, avail_h / pic.height, 1.0 if pic.width < W else 10)
        pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
        pic.left = int((W - pic.width) / 2)
        if sl.get("caption"):
            text(s, M, H - Inches(0.6), W - 2 * M, Inches(0.4), sl["caption"], 12, color=grey)
    elif "table" in sl:
        rows = sl["table"]
        nr, nc = len(rows), max(len(r) for r in rows)
        row_h = Inches(0.75 if nr <= 6 else 0.5 if nr <= 12 else 0.35)
        t = s.shapes.add_table(nr, nc, M, top, W - 2 * M, Emu(min(avail_h, row_h * nr))).table
        for i, r in enumerate(rows):
            for j in range(nc):
                c = t.cell(i, j)
                c.text = str(r[j]) if j < len(r) else ""
                f = c.text_frame.paragraphs[0].runs[0].font if c.text else None
                if f:
                    f.size, f.name, f.bold = Pt(20 if nr <= 6 else 16 if nr <= 12 else 12), FONT, i == 0
                    f.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else dark
                if i == 0:
                    c.fill.solid()
                    c.fill.fore_color.rgb = accent
    elif "columns" in sl:
        cw = (W - 2 * M - Inches(0.5)) / 2
        for k, col in enumerate(sl["columns"][:2]):
            x = M + k * (cw + Inches(0.5))
            text(s, x, top, cw, Inches(0.6), col[0], 22, bold=True, color=accent)
            bullets(s, col[1:], x, top + Inches(0.7), cw, avail_h - Inches(0.7), 18)
    elif "bullets" in sl:
        bullets(s, sl["bullets"], M, top, W - 2 * M, avail_h)
    notes(s, sl.get("notes"))

prs.save(a.out)
print(f"wrote {a.out} ({len(prs.slides)} slides)")
if a.pdf:
    outdir = str(pathlib.Path(a.out).resolve().parent)
    subprocess.run(["fda-soffice", "--convert-to", "pdf", "--outdir", outdir, a.out], check=True,
                   stdout=subprocess.DEVNULL, env={**os.environ})
    print(f"wrote {pathlib.Path(a.out).with_suffix('.pdf')}")
